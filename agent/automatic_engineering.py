from pathlib import Path
import json
import string

from openpyxl import load_workbook
import pandas as pd

MAX_PROFILE_ROWS = 10_000
ROW_SAMPLE = 100


# ---------- TABLE DETECTION (openpyxl) ----------

def get_non_empty_cells(ws):
    """Return a list of (row, col, value) for non-empty cells."""
    cells = []
    for row in ws.iter_rows(values_only=False):  # keep cell objects
        for cell in row:
            if cell.value is not None and str(cell.value).strip() != "":
                cells.append((cell.row, cell.column, cell.value))
    return cells


def detect_table_blocks(ws, min_rows=3, min_cols=2):
    """
    Simple heuristic:
    - find min/max row/col of all non-empty cells
    - treat the whole non-empty rectangle as one table.
    You can later refine this to split on big gaps, headers, etc.
    """
    cells = get_non_empty_cells(ws)
    if not cells:
        return []

    rows = [r for (r, c, v) in cells]
    cols = [c for (r, c, v) in cells]
    min_row, max_row = min(rows), max(rows)
    min_col, max_col = min(cols), max(cols)

    if (max_row - min_row + 1) < min_rows or (max_col - min_col + 1) < min_cols:
        return []

    return [(min_row, max_row, min_col, max_col)]


def block_to_dataframe(ws, block):
    """Convert a block (min_row, max_row, min_col, max_col) into a pandas DataFrame."""
    min_row, max_row, min_col, max_col = block
    data = []
    for row in ws.iter_rows(min_row=min_row, max_row=max_row,
                            min_col=min_col, max_col=max_col,
                            values_only=True):
        data.append(row)

    # Assume first row is header
    header = data[0]
    rows = data[1:]
    df = pd.DataFrame(rows, columns=header)
    return df


def extract_tables_and_images(xlsx_path: str):
    """
    Use openpyxl to:
    - detect table blocks per sheet
    - extract them as DataFrames
    - list images with approximate positions
    """
    wb = load_workbook(xlsx_path, data_only=True)

    tables = []
    images = []

    for ws in wb.worksheets:
        # Tables
        blocks = detect_table_blocks(ws)
        for idx, block in enumerate(blocks):
            df = block_to_dataframe(ws, block)
            table_id = f"{ws.title}_block_{idx}"
            tables.append({
                "table_id": table_id,
                "sheet_name": ws.title,
                "block": {
                    "min_row": block[0],
                    "max_row": block[1],
                    "min_col": block[2],
                    "max_col": block[3],
                },
                "dataframe": df,
            })

        # Images (positions only; you can add file saving later)
        for img_idx, img in enumerate(getattr(ws, "_images", [])):
            row = img.anchor._from.row + 1
            col_idx = img.anchor._from.col
            col_letter = string.ascii_uppercase[col_idx]
            images.append({
                "sheet_name": ws.title,
                "image_id": f"{ws.title}_img_{img_idx}",
                "row": row,
                "col": col_letter,
                # you can add "file_path" or "binary_ref" if you extract/save the image
            })

    return tables, images


# ---------- PROFILING (pandas) ----------

def profile_dataframe(df: pd.DataFrame) -> dict:
    """
    Return a JSON-serializable profile:
    - shape
    - per-column dtypes, missingness, cardinality
    - describe(include='all') stats on a capped sample
    - small sample of rows
    """
    n_rows, n_cols = df.shape

    # Cap rows for profiling heavy stats
    if n_rows > MAX_PROFILE_ROWS:
        df_profile = df.sample(MAX_PROFILE_ROWS, random_state=42)
    else:
        df_profile = df

    dtypes = df_profile.dtypes.astype(str).to_dict()          # column dtypes[web:46]
    missing_count = df_profile.isna().sum().to_dict()         # missing per column
    missing_pct = {
        col: (missing_count[col] / float(len(df_profile))) if len(df_profile) > 0 else 0.0
        for col in df_profile.columns
    }
    nunique = df_profile.nunique(dropna=True).to_dict()       # cardinality[web:36]

    # describe(include='all') for numeric + categorical stats[web:34][web:40][web:47]
    describe_df = df_profile.describe(include="all").transpose()
    describe = describe_df.to_dict(orient="index")

    columns_profile = {}
    for col in df_profile.columns:
        columns_profile[col] = {
            "dtype": dtypes.get(col),
            "missing_count": int(missing_count.get(col, 0)),
            "missing_pct": float(missing_pct.get(col, 0.0)),
            "nunique": int(nunique.get(col, 0)),
            "stats": describe.get(col, {}),
        }

    sample_df = df.head(ROW_SAMPLE)
    sample_records = sample_df.to_dict(orient="records")

    return {
        "n_rows": int(n_rows),
        "n_cols": int(n_cols),
        "columns": list(df.columns),
        "columns_profile": columns_profile,
        "sample": sample_records,
    }


# ---------- MAIN: unified workbook profile ----------

def build_finance_workbook_profile(xlsx_path: str) -> dict:
    """
    Full pipeline:
    - openpyxl: detect tables + images
    - pandas: profile each table
    - return a JSON-ready workbook_profile the agent/LLM can consume
    """
    xlsx_path = Path(xlsx_path)

    tables, images = extract_tables_and_images(str(xlsx_path))

    # Build table profiles
    table_profiles = []
    for t in tables:
        prof = profile_dataframe(t["dataframe"])
        table_profiles.append({
            "table_id": t["table_id"],
            "sheet_name": t["sheet_name"],
            "block": t["block"],
            **prof,
        })

    workbook_profile = {
        "file_name": xlsx_path.name,
        "tables": table_profiles,
        "images": images,
    }

    return workbook_profile


if __name__ == "__main__":
    profile = build_finance_workbook_profile("finance_weird.xlsx")
    print(json.dumps(profile, indent=2))
